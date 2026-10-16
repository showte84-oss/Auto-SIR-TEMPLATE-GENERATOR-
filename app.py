import streamlit as st
import openpyxl
from openpyxl.utils import get_column_letter
import json
import io
import os
import re

st.set_page_config(page_title="SGS SIR Report Automation", page_icon="📊", layout="wide")

hide_streamlit_style = """
<style>
#MainMenu {visibility: hidden;}
header {visibility: hidden;}
footer {visibility: hidden;}
.stDeployButton {display: none;}
[data-testid="stToolbar"] {visibility: hidden !important;}
.watermark {
    position: fixed;
    bottom: 12px;
    left: 50%;
    transform: translateX(-50%);
    opacity: 0.45;
    font-size: 14px;
    font-weight: 700;
    color: #475569;
    letter-spacing: 2px;
    pointer-events: none;
    z-index: 9999;
}
</style>
<div class="watermark">By KNA</div>
"""
st.markdown(hide_streamlit_style, unsafe_allow_html=True)

st.title("📊 SGS Inspection Report (SIR) Excel Generator")
st.markdown("""
ဤ Web App သည် **API Key မလိုဘဲ** SGS Report ပုံများမှ အချက်အလက်များကို `(40 FCL)SIR TEMPLATE.xlsx` ထဲသို့ အလိုအလျောက် ဖြည့်သွင်းပေးမည့် Hybrid စနစ် ဖြစ်ပါသည်။
""")

with st.expander("📌 အဆင့် (၁) - AI တွင် ပုံများဖတ်ရန် Prompt ယူပါ (နှိပ်၍ဖတ်ပါ)", expanded=True):
    st.markdown("""
    ၁။ အခမဲ့ AI တစ်ခုခု (**Gemini Chat** သို့မဟုတ် **ChatGPT**) ကို ဖွင့်ပါ။  
    ၂။ SGS Report ပုံ (၅) ပုံကို တင်ပြီး အောက်ပါ **Prompt** ကို Copy ကူး၍ ပေးပို့လိုက်ပါ။
    """)
    
    ai_prompt = """Read the attached SGS inspection report images (Stuffing Report, Weight List, Tally Form, VGM Field Report, and Container Survey Report) and extract all data into the following exact JSON format. Return ONLY valid JSON inside a ```json ``` codeblock without any extra text.

{
  "job_ref": "2108339",
  "date": "14/10/2026",
  "start_time": "08:30",
  "shipper": "Yangon Agri Food",
  "buyer": "Kalp Impex",
  "commodity": "Black Eye Beans (BEB)",
  "quantity": "48 MTS (1920 Bags)",
  "packing": "New Single pp Bag",
  "poi": "No. 122, HTY.1",
  "survey_location": "No. 122, HTY.1",
  "vessel": "MV As Susanna",
  "destination": "Valencia, Spain",
  "inspector_names": "Kyaw Nyunt, Myint Ko Ko",
  "test_weight": "10 Kgs",
  "test_weight_no": "014",
  "weighing_percent": 0.1,
  "per_bag_gross": 25.05,
  "per_bag_tare": 0.05,
  "weighing_data": [
    {"page": 1, "bags": 192, "gross_weight": 4809.6}
  ],
  "vgm_other_headers": ["Seal", "SECCO", "Paper"],
  "containers": [
    {
      "no": "CAIU 6562141",
      "condition": "Sound",
      "commence": "08:30",
      "complete": "10:00",
      "bags": 960,
      "gross_mt": 24.048,
      "tare_mt": 0.048,
      "seal": "444072",
      "slings": [100, 100, 100, 100, 100, 100, 100, 100, 100, 60],
      "vgm_others": [0.012, 4.0, 10.0],
      "ctn_tare_kg": 2150,
      "hole_dented": "No",
      "floor_type": "Wood",
      "floor_damaged": "No",
      "free_odour": "Yes",
      "free_infestation": "Yes",
      "clean_dry": "Yes",
      "color": "Red",
      "size": "20'",
      "type": "Steel",
      "suitable": "Yes"
    }
  ]
}

Important Rules for Extraction:
1. Container Number (`no`): Always format as 4 English letters, followed by a SINGLE SPACE, followed by the digits (e.g., if handwritten as "BMOU- 1657789" or "CSNU.2702460", output strictly as "BMOU 1657789" or "CSNU 2702460").
2. Container Seal Number (`seal`): Seal numbers contain DIGITS ONLY (0-9). If a handwritten character looks like 'S', interpret it as '5'. If 'O', interpret as '0'. If 'I' or 'l', interpret as '1'. If 'B', interpret as '8'. Output digits only.
3. Format Date as DD/MM/YYYY.
4. Survey fields (hole_dented, floor_type, color, type, etc.) must be in Title Case (First letter Capitalized, e.g., "No", "Wood", "Yes", "Red", "Blue", "Steel").
5. `slings` array can have up to 10 numbers representing each sling column in Tally Form.
6. `vgm_other_headers` must match the exact 3 column headers under 'Other Packages (Kgs)' in the VGM image, and `vgm_others` must contain the 3 corresponding weights."""
    
    st.code(ai_prompt, language="text")

st.divider()

st.subheader("📌 အဆင့် (၂) - Template၊ JSON Data နှင့် File Name ထည့်သွင်းပါ")

if "custom_filename_input" not in st.session_state:
    st.session_state["custom_filename_input"] = "SIR JOB NO"

def update_filename_from_json():
    raw = st.session_state.get("json_input_area", "")
    match = re.search(r'"job_ref"\s*:\s*"([^"]+)"', raw)
    if match and match.group(1).strip():
        st.session_state["custom_filename_input"] = f"SIR JOB NO {match.group(1).strip()}"

col1, col2 = st.columns([1, 2])

with col1:
    default_template = "(40 FCL)SIR TEMPLATE.xlsx"
    use_default = os.path.exists(default_template)
    
    if use_default:
        st.success(f"✅ `{default_template}` ဖိုင်ကို Folder ထဲတွင် အသင့်တွေ့ရှိပါသည်။")
    
    uploaded_template = st.file_uploader(
        "Template ဖိုင်တင်ရန် (Folder ထဲတွင် ရှိပြီးသားဆိုလျှင် ထပ်တင်ရန်မလိုပါ)", 
        type=["xlsx"]
    )

with col2:
    json_input = st.text_area(
        "AI မှ ရရှိလာသော JSON Code ကို ဤနေရာတွင် Paste ချပါ -",
        height=240,
        key="json_input_area",
        on_change=update_filename_from_json,
        placeholder='{"job_ref": "2108339", ...}'
    )
    
    custom_name = st.text_input(
        "✏️ File Name ကို Generate မလုပ်ခင် စိတ်ကြိုက်ပြင်ဆင်နိုင်ပါသည် (`.xlsx` ထည့်ရန်မလိုပါ၊ Auto ထည့်ပေးပါမည်) -",
        key="custom_filename_input",
        placeholder="SIR JOB NO 2108339"
    )

def format_container_no(val):
    if not val:
        return ""
    cleaned = str(val).upper().strip()
    m = re.match(r'^([A-Z]+)[^A-Z0-9]*([0-9]+)$', cleaned)
    if m:
        return f"{m.group(1)} {m.group(2)}"
    return re.sub(r'[^A-Z0-9]+', ' ', cleaned).strip()

def clean_seal_no(val):
    if val is None:
        return ""
    trans = str.maketrans({'S': '5', 's': '5', 'O': '0', 'o': '0', 'I': '1', 'l': '1', 'i': '1', 'B': '8', 'Z': '2', 'z': '2', 'G': '6', 'b': '6'})
    res = str(val).translate(trans)
    digits = re.sub(r'[^0-9]', '', res)
    return digits if digits else res.strip()

def set_cell_value(ws, cell_coord, value):
    cell = ws[cell_coord]
    if isinstance(cell, openpyxl.cell.cell.MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if cell_coord in merged_range:
                top_left_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                top_left_cell.value = value
                return
    else:
        cell.value = value

def generate_excel(template_stream, data):
    wb = openpyxl.load_workbook(template_stream)
    
    job_ref = str(data.get("job_ref", "")).strip()
    date_slash = data.get("date", "")
    start_time = data.get("start_time", "")
    shipper = data.get("shipper", "")
    buyer = data.get("buyer", "")
    commodity = data.get("commodity", "")
    quantity = data.get("quantity", "")
    packing = data.get("packing", "")
    poi = data.get("poi", "")
    survey_loc = data.get("survey_location", poi)
    vessel = data.get("vessel", "")
    destination = data.get("destination", "")
    inspector_names = data.get("inspector_names", "")
    test_weight = data.get("test_weight", "10 Kgs")
    test_weight_no = data.get("test_weight_no", "")
    weighing_percent = data.get("weighing_percent", 0.1)
    per_bag_gross = data.get("per_bag_gross", 0)
    per_bag_tare = data.get("per_bag_tare", 0)
    weighing_data = data.get("weighing_data", [])
    vgm_headers = data.get("vgm_other_headers", ["Seal", "SECCO", "Paper"])
    containers = data.get("containers", [])
    visible_count = max(10, len(containers))
    
    # 1. STUFFING REPORT (MASTER PAGE)
    ws_stuffing = wb['STUFFING REPORT']
    set_cell_value(ws_stuffing, 'E5', job_ref)
    set_cell_value(ws_stuffing, 'F6', date_slash)
    set_cell_value(ws_stuffing, 'F7', date_slash)
    set_cell_value(ws_stuffing, 'E8', shipper)
    set_cell_value(ws_stuffing, 'E9', buyer)
    set_cell_value(ws_stuffing, 'E10', commodity)
    set_cell_value(ws_stuffing, 'E11', quantity)
    set_cell_value(ws_stuffing, 'E12', packing)
    set_cell_value(ws_stuffing, 'E13', poi)
    set_cell_value(ws_stuffing, 'E14', vessel)
    set_cell_value(ws_stuffing, 'E15', destination)
    set_cell_value(ws_stuffing, 'E16', inspector_names)
    
    start_row = 21
    for i, c in enumerate(containers):
        row = start_row + i
        set_cell_value(ws_stuffing, f'B{row}', format_container_no(c.get("no", "")))
        set_cell_value(ws_stuffing, f'C{row}', c.get("condition", "SOUND"))
        set_cell_value(ws_stuffing, f'D{row}', c.get("commence", ""))
        set_cell_value(ws_stuffing, f'E{row}', c.get("complete", ""))
        set_cell_value(ws_stuffing, f'F{row}', c.get("bags", 0))
        set_cell_value(ws_stuffing, f'G{row}', c.get("gross_mt", 0))
        set_cell_value(ws_stuffing, f'H{row}', c.get("tare_mt", 0))
        set_cell_value(ws_stuffing, f'I{row}', clean_seal_no(c.get("seal", "")))
        
    for r in range(start_row + len(containers), 61):
        for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
            set_cell_value(ws_stuffing, f'{col}{r}', None)
            
    for r in range(start_row, 61):
        ws_stuffing.row_dimensions[r].hidden = (r >= start_row + visible_count)
        
    set_cell_value(ws_stuffing, 'C68', test_weight)
    set_cell_value(ws_stuffing, 'G68', test_weight_no)
    
    # 2. WEIGHT LIST (SUMMARY)
    ws_weight = wb['WEIGHT LIST (SUMMARY)']
    set_cell_value(ws_weight, 'C8', weighing_percent)
    
    for i, d in enumerate(weighing_data):
        row = 12 + i
        set_cell_value(ws_weight, f'A{row}', d.get("page", i+1))
        set_cell_value(ws_weight, f'B{row}', d.get("bags", 0))
        set_cell_value(ws_weight, f'C{row}', d.get("gross_weight", 0))
        set_cell_value(ws_weight, f'D{row}', "kgs")
        
    for r in range(12 + len(weighing_data), 27):
        for col in ['A', 'B', 'C', 'D']:
            set_cell_value(ws_weight, f'{col}{r}', None)
            
    set_cell_value(ws_weight, 'G13', per_bag_tare)
    
    # 3. TALLY FORM (Preserve Column B Master Link & Column M Formula)
    ws_tally = wb['TALLY FORM']
    start_row_tally = 11
    for i, c in enumerate(containers):
        row = start_row_tally + i
        slings = c.get("slings", [])
        for col_idx in range(10): # C to L only
            col_letter = get_column_letter(3 + col_idx)
            val = slings[col_idx] if col_idx < len(slings) else None
            set_cell_value(ws_tally, f'{col_letter}{row}', val)
            
    for r in range(start_row_tally + len(containers), 51):
        for col in range(3, 13): # Clear C to L only; keep B & M formulas intact
            set_cell_value(ws_tally, f'{get_column_letter(col)}{r}', None)
            
    for r in range(start_row_tally, 51):
        ws_tally.row_dimensions[r].hidden = (r >= start_row_tally + visible_count)
        
    # 4. VGM REPORT (Preserve Column B Master Link & E, L Formulas)
    ws_vgm = wb['VGM REPORT']
    set_cell_value(ws_vgm, 'J5', date_slash)
    set_cell_value(ws_vgm, 'L5', start_time)
    
    for idx, h_Col in enumerate(['F9', 'G9', 'H9']):
        if idx < len(vgm_headers):
            set_cell_value(ws_vgm, h_Col, vgm_headers[idx])
            
    start_row_vgm = 10
    for i, c in enumerate(containers):
        row = start_row_vgm + i
        set_cell_value(ws_vgm, f'C{row}', c.get("bags", 0))
        set_cell_value(ws_vgm, f'D{row}', per_bag_gross)
        vgm_others = c.get("vgm_others", [0, 0, 0])
        set_cell_value(ws_vgm, f'F{row}', vgm_others[0] if len(vgm_others) > 0 else None)
        set_cell_value(ws_vgm, f'G{row}', vgm_others[1] if len(vgm_others) > 1 else None)
        set_cell_value(ws_vgm, f'H{row}', vgm_others[2] if len(vgm_others) > 2 else None)
        set_cell_value(ws_vgm, f'K{row}', c.get("ctn_tare_kg", 0))
        
    for r in range(start_row_vgm + len(containers), 50):
        for col in ['C', 'D', 'F', 'G', 'H', 'I', 'J', 'K']: # Keep B, E, L formulas intact
            set_cell_value(ws_vgm, f'{col}{r}', None)
            
    for r in range(start_row_vgm, 50):
        ws_vgm.row_dimensions[r].hidden = (r >= start_row_vgm + visible_count)
        
    # 5. CTN SURVEY REPORT (Preserve Column B Master Link)
    ws_survey = wb['CTN SURVEY REPORT']
    set_cell_value(ws_survey, 'I5', survey_loc)
    set_cell_value(ws_survey, 'I6', date_slash)
    
    start_row_survey = 9
    for i, c in enumerate(containers):
        row = start_row_survey + i
        set_cell_value(ws_survey, f'C{row}', str(c.get("hole_dented", "No")).title())
        set_cell_value(ws_survey, f'D{row}', str(c.get("floor_type", "Wood")).title())
        set_cell_value(ws_survey, f'E{row}', str(c.get("floor_damaged", "No")).title())
        set_cell_value(ws_survey, f'F{row}', str(c.get("free_odour", "Yes")).title())
        set_cell_value(ws_survey, f'G{row}', str(c.get("free_infestation", "Yes")).title())
        set_cell_value(ws_survey, f'H{row}', str(c.get("clean_dry", "Yes")).title())
        set_cell_value(ws_survey, f'I{row}', str(c.get("color", "")).title())
        set_cell_value(ws_survey, f'J{row}', c.get("size", "20'"))
        set_cell_value(ws_survey, f'K{row}', str(c.get("type", "Steel")).title())
        set_cell_value(ws_survey, f'L{row}', str(c.get("suitable", "Yes")).title())
        
    for r in range(start_row_survey + len(containers), 49):
        for col in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L']: # Keep Column B formula intact
            set_cell_value(ws_survey, f'{col}{r}', None)
            
    for r in range(start_row_survey, 49):
        ws_survey.row_dimensions[r].hidden = (r >= start_row_survey + visible_count)
        
    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream.getvalue(), job_ref

if st.button("🚀 Generate Excel File", type="primary"):
    if not json_input.strip():
        st.error("⚠️ ကျေးဇူးပြု၍ JSON Data ကို အရင်ထည့်သွင်းပေးပါ။")
    else:
        try:
            cleaned_json = json_input.strip()
            if cleaned_json.startswith("```"):
                cleaned_json = cleaned_json.split("\n", 1)[1]
                if cleaned_json.endswith("```"):
                    cleaned_json = cleaned_json.rsplit("```", 1)[0]
            
            parsed_data = json.loads(cleaned_json)
            template_source = uploaded_template if uploaded_template else (default_template if use_default else None)
            
            if not template_source:
                st.error("⚠️ `(40 FCL)SIR TEMPLATE.xlsx` ဖိုင်ကို Upload တင်ပေးပါ။")
            else:
                excel_bytes, job_ref = generate_excel(template_source, parsed_data)
                
                clean_name = custom_name.strip()
                if not clean_name or clean_name == "SIR JOB NO":
                    clean_name = f"SIR JOB NO {job_ref}" if job_ref else "SIR JOB NO"
                
                if clean_name.lower().endswith(".xlsx"):
                    final_filename = clean_name
                else:
                    final_filename = f"{clean_name}.xlsx"
                
                st.session_state['excel_bytes'] = excel_bytes
                st.session_state['final_filename'] = final_filename
                st.success(f"🎉 `{final_filename}` ဖိုင်ကို အောင်မြင်စွာ ဖန်တီးပြီးပါပြီ! အောက်ပါခလုတ်ကိုနှိပ်၍ Download ရယူပါ။")
        except json.JSONDecodeError as e:
            st.error(f"❌ JSON Format မှားယွင်းနေပါသည်: {e}")
        except Exception as e:
            st.error(f"❌ အမှားအယွင်းတစ်ခု ဖြစ်ပေါ်နေပါသည်: {e}")

if 'excel_bytes' in st.session_state:
    st.download_button(
        label=f"📥 Download ({st.session_state['final_filename']})",
        data=st.session_state['excel_bytes'],
        file_name=st.session_state['final_filename'],
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )
